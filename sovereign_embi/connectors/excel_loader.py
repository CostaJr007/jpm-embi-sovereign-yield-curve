"""Excel Loader & Curve Data Connector.

Extracts sovereign bond yields, EMBI+ spreads, and macro variables
from Excel workbooks ('EMBI DOS DEUSES ATUALIZANDO REAL TIME.xlsx')
and provides synthetic fallback datasets for offline execution.
"""

from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional
import openpyxl
import numpy as np


@dataclass
class SovereignDataset:
    """Complete cross-border sovereign dataset."""
    br_maturities: List[float]
    br_yields: List[float]
    us_maturities: List[float]
    us_yields: List[float]
    embi_bps: float
    selic_rate: float
    fed_funds_rate: float
    br_inflation_ipca: float
    us_inflation_cpi: float
    br_gdp_forecast: float
    us_gdp_forecast: float


class SovereignDataLoader:
    """Parses Excel workbooks or generates synthetic yield curve datasets."""

    TENOR_MAP = {
        "3 meses": 0.25,
        "6 meses": 0.50,
        "9 meses": 0.75,
        "1 ano": 1.0,
        "2 anos": 2.0,
        "3 anos": 3.0,
        "5 anos": 5.0,
        "7 anos": 7.0,
        "8 anos": 8.0,
        "10 anos": 10.0,
        "30 anos": 30.0
    }

    @classmethod
    def load_from_excel(cls, file_path: str) -> SovereignDataset:
        """Parse 'EMBI DOS DEUSES ATUALIZANDO REAL TIME.xlsx' workbook."""
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"Workbook not found at {file_path}")

        wb = openpyxl.load_workbook(str(path), data_only=True)
        sheet_comp = wb["comparativo de retorno"]

        # Parse BR Yields from rows 3 to 11
        br_maturities = []
        br_yields = []
        for r in range(3, 12):
            label = sheet_comp.cell(r, 2).value
            yield_val = sheet_comp.cell(r, 3).value
            if label and yield_val is not None:
                tenor_str = str(label).replace("Brasil a ", "").strip().lower()
                tenor = cls.TENOR_MAP.get(tenor_str, None)
                if tenor is not None:
                    br_maturities.append(float(tenor))
                    br_yields.append(float(yield_val))

        # Parse US Yields from rows 18 to 26
        us_maturities = []
        us_yields = []
        for r in range(18, 27):
            label = sheet_comp.cell(r, 2).value
            yield_val = sheet_comp.cell(r, 3).value
            if label and yield_val is not None:
                tenor_str = str(label).replace("EUA a ", "").strip().lower()
                tenor = cls.TENOR_MAP.get(tenor_str, None)
                if tenor is not None:
                    us_maturities.append(float(tenor))
                    us_yields.append(float(yield_val))

        # If US yields not in sheet_comp, check Planilha1
        if len(us_maturities) == 0 and "Planilha1" in wb.sheetnames:
            p1 = wb["Planilha1"]
            for r in range(3, 12):
                label = p1.cell(r, 3).value
                yield_val = p1.cell(r, 4).value
                if label and yield_val is not None:
                    tenor_str = str(label).replace("EUA a ", "").strip().lower()
                    tenor = cls.TENOR_MAP.get(tenor_str, None)
                    if tenor is not None:
                        us_maturities.append(float(tenor))
                        us_yields.append(float(yield_val))

        # Parse Macro constants
        # Cell W2: EMBI BR (e.g. 334 bps)
        embi_bps = 334.0
        try:
            val = sheet_comp.cell(2, 23).value  # W2
            if val is not None:
                embi_bps = float(val)
        except Exception:
            pass

        # Cell R2: Selic (0.1075 -> 10.75)
        selic_rate = 10.75
        try:
            val = sheet_comp.cell(2, 18).value  # R2
            if val is not None:
                selic_rate = float(val) * 100.0 if float(val) < 1.0 else float(val)
        except Exception:
            pass

        # Cell M2: Fed Funds (0.0025 -> 0.25)
        fed_funds = 0.25
        try:
            val = sheet_comp.cell(2, 13).value  # M2
            if val is not None:
                fed_funds = float(val) * 100.0 if float(val) < 1.0 else float(val)
        except Exception:
            pass

        # Cell S2: IPCA (0.1054 -> 10.54)
        ipca = 10.54
        try:
            val = sheet_comp.cell(2, 19).value  # S2
            if val is not None:
                ipca = float(val) * 100.0 if float(val) < 1.0 else float(val)
        except Exception:
            pass

        # Cell N2: CPI (0.079 -> 7.90)
        cpi = 7.90
        try:
            val = sheet_comp.cell(2, 14).value  # N2
            if val is not None:
                cpi = float(val) * 100.0 if float(val) < 1.0 else float(val)
        except Exception:
            pass

        return SovereignDataset(
            br_maturities=br_maturities,
            br_yields=br_yields,
            us_maturities=us_maturities,
            us_yields=us_yields,
            embi_bps=embi_bps,
            selic_rate=selic_rate,
            fed_funds_rate=fed_funds,
            br_inflation_ipca=ipca,
            us_inflation_cpi=cpi,
            br_gdp_forecast=3.4,
            us_gdp_forecast=7.0
        )

    @classmethod
    def generate_sample_dataset(cls) -> SovereignDataset:
        """Generate realistic baseline Brazil vs US yield curve dataset."""
        return SovereignDataset(
            br_maturities=[0.25, 0.50, 0.75, 1.0, 2.0, 3.0, 5.0, 8.0, 10.0],
            br_yields=[12.245, 12.939, 13.422, 13.494, 12.556, 12.690, 12.286, 12.196, 12.207],
            us_maturities=[0.25, 0.50, 1.0, 2.0, 3.0, 5.0, 7.0, 10.0, 30.0],
            us_yields=[0.495, 0.871, 1.268, 1.940, 2.129, 2.350, 2.450, 2.520, 2.850],
            embi_bps=334.0,
            selic_rate=10.75,
            fed_funds_rate=0.25,
            br_inflation_ipca=10.54,
            us_inflation_cpi=7.90,
            br_gdp_forecast=3.4,
            us_gdp_forecast=7.0
        )
